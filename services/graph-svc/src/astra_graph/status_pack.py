"""The Status Pack -- story S10.2.1, opening F10.2.

    "Status Pack: generated weekly as an editable narrative with the numbers and charts
    of the Board, exportable to PDF and PPTX; edits are stored with the version"

§15.3.1's own row, verbatim: *"Generated weekly pack: progress, gates, exceptions
aging, risks; editable narrative | Generate; edit; publish to client."*

**The narrative is composed from the identical real figures the Programme Board itself
already computes** (`programme_surface.kpi_strip`/`train_swimlanes`/`milestone_rail`,
`exception_ageing.exception_ageing`) -- "the status pack writes itself" (this story's
own "So that") means assembling real numbers into real sentences, not inventing a
second source of truth for any of them. "Risks" (the spec's own word, absent from the
backlog AC) is read as the same real blocked-reasons and breached-SLA facts the KPI
strip and swimlanes already carry -- there is no separate risk register anywhere in
this codebase to draw from instead.

**"Generated weekly" is a `POST`-triggered snapshot, not a background scheduler.** The
identical disclosed posture `g2_reminders.send_due_reminders` already set: safe to call
repeatedly (a fresh call for a week that already has a pack starts a new version rather
than erasing the old one), real and recorded the moment someone asks for it, with a
background loop that calls it on a timer left as real future scope this story does not
claim (the same `HarvestScheduler`-shaped gap `g2_reminders.py`'s own docstring already
names for reminders).

**Every edit is a new row, never an in-place update -- "edits are stored with the
version" is the AC's own literal words**, and migration v0037's own docstring gives the
identical reasoning `calibration_baseline` already has for its own append-only shape.

**"Publish to client" records a real state transition (`published_at`), not a real
delivery.** No outward delivery channel exists anywhere in this codebase (confirmed:
the identical disclosed gap `g2_reminders.py`'s own "sent" already carries) -- the same
"a real record, not a claim of delivery nobody could verify" posture this codebase
takes throughout.

**PDF and PPTX are real, generated documents -- text and tables, not native charts.**
`reportlab` (PDF) and `python-pptx` (PPTX) are both real, working, newly-added
dependencies (see `pyproject.toml`'s own comment); "the numbers... of the Board" are
rendered as real tables carrying the identical figures the console itself shows.
Rendering those same figures as native bar/line charts in either format is real,
separate future work this story does not take on -- disclosed here rather than half
attempted.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from typing import Any

import asyncpg
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .adoption import AdoptionConfigStore, AdoptionStore
from .errors import InvalidRequestError
from .exception_ageing import exception_ageing
from .g2 import QuestionStore
from .ids import new_ulid
from .invoicing import UnitPriceStore
from .programme_surface import kpi_strip, milestone_rail, train_swimlanes

PACK_TABLE = "public.status_pack"


@dataclass(frozen=True, slots=True)
class StatusPack:
    id: str
    week_of: str
    version: int
    narrative: str
    report: dict[str, Any]
    generated_by: str
    generated_at: str
    published_at: str | None

    def as_dict(self) -> dict[str, Any]:
        return {
            "id": self.id, "week_of": self.week_of, "version": self.version,
            "narrative": self.narrative, "report": self.report,
            "generated_by": self.generated_by, "generated_at": self.generated_at,
            "published_at": self.published_at,
        }


def _week_of(today: date | None = None) -> date:
    """The Monday of the current ISO week -- one pack per calendar week, matching the
    AC's own "generated weekly."""
    day = today or datetime.now(UTC).date()
    return day - timedelta(days=day.weekday())


def _narrative(report: dict[str, Any]) -> str:
    kpis = report["kpis"]
    swimlanes = report["swimlanes"]
    ageing = report["exception_ageing"]

    mus_by_state = ", ".join(f"{count} {state}" for state, count in kpis["mus_by_state"].items()) or "none yet"
    first_pass = kpis["first_pass_parity"]["first_pass_rate"]
    first_pass_text = f"{first_pass * 100:.1f}%" if first_pass is not None else "not yet measured"
    absorption = kpis["absorption"]["mean_ratio"]
    absorption_text = f"{absorption * 100:.1f}%" if absorption is not None else "not yet measured"
    spend = kpis["spend_vs_budget"]
    blocked_total = sum(t["blocked_count"] for t in swimlanes["trains"])

    lines = [
        f"Migration Units by state: {mus_by_state}.",
        f"First-pass parity across the live estate: {first_pass_text}.",
        f"Absorption against the configured threshold ({kpis['absorption']['threshold'] * 100:.0f}%): "
        f"{absorption_text}.",
        f"Gates due this week: {kpis['gates_due_this_week']['due_this_week_count']} "
        f"({kpis['gates_due_this_week']['already_breached_count']} already past SLA; {kpis['gates_due_this_week']['scope']}).",
        f"Spend to date: ${spend['spend']:,.0f} against a planned value of ${spend['budget']:,.0f} "
        f"({'ahead of' if spend['delta'] >= 0 else 'behind'} plan by ${abs(spend['delta']):,.0f}).",
        f"{blocked_total} Migration Unit(s) are currently blocked across {len(swimlanes['trains'])} train(s).",
        f"{ageing.get('total_open', 0)} exception(s) are open or blocked; "
        + (
            f"the Mender closes {ageing['mender_close_rate']['rate'] * 100:.0f}% of failures without a decision."
            if ageing.get("mender_close_rate", {}).get("rate") is not None
            else "no failure has ever opened an ExceptionCase yet."
        ),
    ]
    return "\n".join(lines)


async def _assemble_report(
    pool: asyncpg.Pool,
    graph_name: str,
    *,
    question_store: QuestionStore,
    unit_price_store: UnitPriceStore,
    adoption_store: AdoptionStore,
    adoption_config_store: AdoptionConfigStore,
) -> dict[str, Any]:
    return {
        "kpis": await kpi_strip(
            pool, graph_name, question_store=question_store, unit_price_store=unit_price_store,
            adoption_store=adoption_store, adoption_config_store=adoption_config_store,
        ),
        "swimlanes": await train_swimlanes(pool, graph_name),
        "milestones": await milestone_rail(pool, graph_name),
        "exception_ageing": await exception_ageing(pool, graph_name),
    }


def _from_row(row: asyncpg.Record) -> StatusPack:
    report = row["report"]
    return StatusPack(
        id=row["id"],
        week_of=row["week_of"].isoformat(),
        version=int(row["version"]),
        narrative=row["narrative"],
        report=report if isinstance(report, dict) else json.loads(report),
        generated_by=row["generated_by"],
        generated_at=row["generated_at"].isoformat(),
        published_at=row["published_at"].isoformat() if row["published_at"] else None,
    )


async def latest_pack(pool: asyncpg.Pool, graph_name: str) -> StatusPack | None:
    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            f"SELECT * FROM {PACK_TABLE} WHERE graph = $1 ORDER BY week_of DESC, version DESC LIMIT 1",
            graph_name,
        )
    return _from_row(row) if row else None


async def generate_pack(
    pool: asyncpg.Pool,
    graph_name: str,
    *,
    question_store: QuestionStore,
    unit_price_store: UnitPriceStore,
    adoption_store: AdoptionStore,
    adoption_config_store: AdoptionConfigStore,
    generated_by: str,
) -> StatusPack:
    report = await _assemble_report(
        pool, graph_name, question_store=question_store, unit_price_store=unit_price_store,
        adoption_store=adoption_store, adoption_config_store=adoption_config_store,
    )
    narrative = _narrative(report)
    week_of = _week_of()
    pack_id = f"statuspack_{new_ulid()}"

    async with pool.acquire() as conn, conn.transaction():
        version = (
            await conn.fetchval(
                f"SELECT MAX(version) FROM {PACK_TABLE} WHERE graph = $1 AND week_of = $2",
                graph_name, week_of,
            )
            or 0
        ) + 1
        row = await conn.fetchrow(
            f"""INSERT INTO {PACK_TABLE}
                 (id, graph, week_of, version, narrative, report, generated_by)
                 VALUES ($1, $2, $3, $4, $5, $6::jsonb, $7)
                 RETURNING *""",
            pack_id, graph_name, week_of, version, narrative, json.dumps(report), generated_by,
        )
    assert row is not None
    return _from_row(row)


async def edit_pack(
    pool: asyncpg.Pool, graph_name: str, *, narrative: str, edited_by: str,
) -> StatusPack:
    """A new version carrying the edited narrative and the same report snapshot the
    pack being edited already had -- an edit changes the words, never the frozen
    numbers a version was generated against."""
    current = await latest_pack(pool, graph_name)
    if current is None:
        raise InvalidRequestError("no Status Pack exists yet to edit -- generate one first")

    week_of = date.fromisoformat(current.week_of)
    pack_id = f"statuspack_{new_ulid()}"
    async with pool.acquire() as conn, conn.transaction():
        version = (
            await conn.fetchval(
                f"SELECT MAX(version) FROM {PACK_TABLE} WHERE graph = $1 AND week_of = $2",
                graph_name, week_of,
            )
            or 0
        ) + 1
        row = await conn.fetchrow(
            f"""INSERT INTO {PACK_TABLE}
                 (id, graph, week_of, version, narrative, report, generated_by)
                 VALUES ($1, $2, $3, $4, $5, $6::jsonb, $7)
                 RETURNING *""",
            pack_id, graph_name, week_of, version, narrative, json.dumps(current.report), edited_by,
        )
    assert row is not None
    return _from_row(row)


async def publish_pack(pool: asyncpg.Pool, graph_name: str) -> StatusPack:
    current = await latest_pack(pool, graph_name)
    if current is None:
        raise InvalidRequestError("no Status Pack exists yet to publish -- generate one first")

    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            f"""UPDATE {PACK_TABLE} SET published_at = now()
                 WHERE graph = $1 AND id = $2
                 RETURNING *""",
            graph_name, current.id,
        )
    assert row is not None
    return _from_row(row)


# ------------------------------------------------------------------------------ export


def render_pdf(pack: StatusPack) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER)
    styles = getSampleStyleSheet()
    story: list[Any] = [
        Paragraph(f"Astra Data — Status Pack, week of {pack.week_of}", styles["Title"]),
        Paragraph(f"Version {pack.version} · generated by {pack.generated_by} at {pack.generated_at}", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Narrative", styles["Heading2"]),
    ]
    for line in pack.narrative.split("\n"):
        story.append(Paragraph(line, styles["Normal"]))
    story.append(Spacer(1, 12))
    story.append(Paragraph("KPI strip", styles["Heading2"]))

    kpis = pack.report["kpis"]
    kpi_rows = [["KPI", "Value"]]
    kpi_rows.append(["MUs by state", ", ".join(f"{k}: {v}" for k, v in kpis["mus_by_state"].items()) or "—"])
    fp = kpis["first_pass_parity"]["first_pass_rate"]
    kpi_rows.append(["First-pass parity", f"{fp * 100:.1f}%" if fp is not None else "—"])
    absorption = kpis["absorption"]["mean_ratio"]
    kpi_rows.append(["Absorption", f"{absorption * 100:.1f}%" if absorption is not None else "—"])
    kpi_rows.append(["Gates due this week", str(kpis["gates_due_this_week"]["due_this_week_count"])])
    spend = kpis["spend_vs_budget"]
    kpi_rows.append(["Spend vs budget", f"${spend['spend']:,.0f} / ${spend['budget']:,.0f}"])

    table = Table(kpi_rows, colWidths=[200, 250])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f5fd0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buffer.getvalue()


def render_pptx(pack: StatusPack) -> bytes:
    presentation = Presentation()
    title_slide_layout = presentation.slide_layouts[0]
    slide = presentation.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = f"Status Pack — week of {pack.week_of}"
    slide.placeholders[1].text = f"Version {pack.version} · {pack.generated_by} · {pack.generated_at}"

    narrative_layout = presentation.slide_layouts[1]
    narrative_slide = presentation.slides.add_slide(narrative_layout)
    narrative_slide.shapes.title.text = "Narrative"
    body = narrative_slide.placeholders[1].text_frame
    lines = pack.narrative.split("\n")
    body.text = lines[0] if lines else ""
    for line in lines[1:]:
        paragraph = body.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(14)

    kpi_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    kpi_slide.shapes.title.text = "KPI strip"
    kpis = pack.report["kpis"]
    rows = [
        ("MUs by state", ", ".join(f"{k}: {v}" for k, v in kpis["mus_by_state"].items()) or "—"),
        ("First-pass parity", _pct(kpis["first_pass_parity"]["first_pass_rate"])),
        ("Absorption", _pct(kpis["absorption"]["mean_ratio"])),
        ("Gates due this week", str(kpis["gates_due_this_week"]["due_this_week_count"])),
        ("Spend vs budget", f"${kpis['spend_vs_budget']['spend']:,.0f} / ${kpis['spend_vs_budget']['budget']:,.0f}"),
    ]
    table_shape = kpi_slide.shapes.add_table(
        len(rows) + 1, 2, Inches(0.5), Inches(1.5), Inches(9), Inches(0.4 * (len(rows) + 1))
    )
    table = table_shape.table
    table.cell(0, 0).text = "KPI"
    table.cell(0, 1).text = "Value"
    for index, (label, value) in enumerate(rows, start=1):
        table.cell(index, 0).text = label
        table.cell(index, 1).text = value

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pct(value: float | None) -> str:
    return f"{value * 100:.1f}%" if value is not None else "—"


__all__ = [
    "PACK_TABLE",
    "StatusPack",
    "edit_pack",
    "generate_pack",
    "latest_pack",
    "publish_pack",
    "render_pdf",
    "render_pptx",
]
